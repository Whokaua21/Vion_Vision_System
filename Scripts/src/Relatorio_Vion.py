import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from PIL import Image
from io import BytesIO
import base64
from docx import Document
from datetime import datetime
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

#Esse arquivo e o Arquivo que tem como funcionalidade criar um relatorio docx do problema analisado pela camera
# Ele criar um relatorio utilizando IA do problema
# Itens que são caracterizados no Sistema: Tipo de Ocorrência, Evidencias, Resumo, Data,Horario e Local, Nivel do Risco,
#  Açoes recomendadas, Confiabilidade da analise e Conclusão 
DATE_TIME = datetime.now()

RELATORIO_AGENTE_IA = f"""
Você é o Agente de Relatórios Inteligentes da Vaion Vision, especializado em analisar imagens de câmeras de segurança e gerar relatórios profissionais sobre possíveis ocorrências.
Sua função é analisar as informações recebidas, identificar evidências de situações como furtos, roubos, invasões, acidentes, quedas e comportamentos perigosos, descrevendo os acontecimentos de maneira objetiva e imparcial.
Nunca invente informações, horários, identidades ou ações que não estejam presentes nas evidências. Diferencie fatos observados de suspeitas e não afirme que alguém cometeu um crime sem confirmação. Caso os dados sejam insuficientes, informe essa limitação.
Gere sempre o relatório no seguinte formato:
    Tipo de ocorrência: 
    Classifique a situação identificada.
    Resumo: 
    Descreva brevemente o que aconteceu.
    Evidências: 
    Liste os elementos e acontecimentos observados.
    Data, horário e local:
    Informe os dados disponíveis.
    Nível de risco:
    Classifique como BAIXO, MODERADO, ALTO ou CRÍTICO e justifique a classificação.
    Ações recomendadas:
    Sugira medidas adequadas para verificar ou resolver a situação.
    Confiabilidade da análise: Informe se as evidências são suficientes, parciais ou inconclusivas.
    Conclusão: Apresente uma síntese da ocorrência e a providência prioritária recomendada.
Respeite o formado dos topicos, não quero nenhum adicional
Utilize linguagem profissional, clara, técnica e direta. Priorize informações relevantes para agilizar a tomada de decisões e auxiliar os responsáveis pela segurança.
Se nenhuma ocorrência relevante puder ser identificada, informe que não existem evidências suficientes para confirmar uma ocorrência. A ausência de detecção não garante que nada aconteceu.
Seu objetivo é transformar informações visuais em relatórios úteis, contribuindo para um monitoramento mais eficiente e uma resposta humana mais rápida e organizada.

informações adcionais: tempo: {DATE_TIME}, 
local: Indefinido, 
Camera: indefinido.

Obs: Quero que voce de uma quebra de linha em cada topico e não coloque ** .
ex: Nivel de risco: (Quebra a linha)Texto
"""
#Nesse parte do sistema estamos usando o diversas bibliotecas
#E necessario futuramente ver a segurança do sistema
# Sugestão: Tem que ter um jeito de vedar a parte da Ia pra não ser algo de facil acesso , e melhorar a parte do style do docx
env_ = "Scripts\src\.env"
load_dotenv(env_)
class Vaion_vision_Relatorio:
    def __init__(self, imagem:str):
        # Atributos para a classe Vaion_vision_Relatorio
        self.imagem_vision = imagem
    def Relatorio_IA(self):
        #Esse codigo e para converter a imagem para base64 para a IA ler, porque a IA le ou usando URL ou base64

        extensao = os.path.splitext(self.imagem_vision)[1]
        if extensao in [".jpeg",".png",".jfif",".jpg"]:
            with Image.open(self.imagem_vision) as image:
                buffer = BytesIO()
                image.save(buffer, format="JPEG")

            image_base64 = base64.b64encode(
                buffer.getvalue()
            ).decode("utf-8")
        else:
            return "Imagem fora do tipo"
        #KEY_API_GROQ_IA e uma variavel que existe no .env do projeto que possui o codigo de acesso a IA
        # Para poder usar a IA voce precisa da chave que esta dentro do arquivo .env 
        KEY_API_GROQ = Groq(
           api_key= os.getenv("KEY_API_GROQ_IA")
        )
        try:
            #Essa parte e a parte de uso da IA, aqui temos o promt para ele fazer oque deve ser feito e a imagem que ele ira analisar utilizando o base64
            RES_GROQ = KEY_API_GROQ.chat.completions.create(
                model="qwen/qwen3.8-27b",
                messages=[
                    {
                        "role":"user",
                        "content":[
                            {
                                "type":"text",
                                "text":RELATORIO_AGENTE_IA
                            },
                            {
                                "type":"image_url",
                                "image_url":{
                                    "url":f"data:image/jpeg;base64,{image_base64}"
                                }
                            }
                        ]
                    }
                ],
                temperature=1,
                top_p=1,
                stream=False,
                stop=None,
                max_completion_tokens=500

            )
        except ConnectionError:
            print("Token esgotado ou modelo de IA invalido")

        #Estilização do documento relatorio do sistema
        date_ = datetime.today().date()
        texto = RES_GROQ.choices[0].message.content

        relatorio_vaion_docx = Document()
        path_docs = os.listdir(r'Scripts\src\docs')       
        
        Title_relatorio = relatorio_vaion_docx.add_paragraph(f"Relatorio {date_} - Vaion Vision")
        Version_relatorio = relatorio_vaion_docx.add_paragraph("Versão 0.2")
        Title_relatorio.alignment = WD_ALIGN_PARAGRAPH.CENTER
        Version_relatorio.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in Title_relatorio.runs:
            run.bold = True
            run.font.size = Pt(20)
            run.font.name = "Franklin Gothic Book"


        for run in Version_relatorio.runs:
            run.bold = True
            run.font.size = Pt(13)
            run.font.name = "Franklin Gothic Book"
            

        partes = texto.split("\n")

        for i in partes:
            if i in ["Tipo de ocorrência:","Resumo:","Evidências:","Data, horário e local:","Nível de risco:","Ações recomendadas:","Confiabilidade da análise:","Conclusão:"]:
                title = relatorio_vaion_docx.add_paragraph(i)
                title.alignment = WD_ALIGN_PARAGRAPH.LEFT

                for run in title.runs:
                    run.bold = True
                    run.font.size = Pt(15)
                    run.font.name = "Franklin Gothic Book"
                    
            else:
                relatorio_vaion_docx.add_paragraph(i)
        data_relatorio = relatorio_vaion_docx.add_paragraph(f"\nRelatorio feito no dia: {DATE_TIME.strftime("%d/%m/%Y %H:%M:%S")}")
        for run in data_relatorio.runs:
            run.bold = True
            run.font.size = Pt(9)
            run.font.name = "Franklin Gothic Book"

        infor_vaion = relatorio_vaion_docx.add_paragraph("VAION VISION | Relatório gerado por Inteligência Artificial\n" \
        "PROTOTIPO | Esse relatorio e um protipo de uma das funções do Vaion Vision, possiveis erros são comuns nessa etapa")

        for run in infor_vaion.runs:
            run.font.size = Pt(9)
            run.font.name = "Franklin Gothic Book"
            run.bold = True
      

        relatorio_vaion_docx.save(f"Scripts\src\docs\Relatorio_{date_}-{path_docs.__len__()}.docx")
        return ":Relatorio feito com Sucesso:"

# Para chamar o sistema dentro da class e necessrio fazer isso, esse sistema por enquanto esta só precisando de uma imagem para funcionar 
Vaion = Vaion_vision_Relatorio(r"Scripts\src\\image_test\\e011.jpg")

print(Vaion.Relatorio_IA())